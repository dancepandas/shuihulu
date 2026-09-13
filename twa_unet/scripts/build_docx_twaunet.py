# -*- coding: utf-8 -*-
"""
《水利水电技术（中英文）》投稿稿生成脚本 —— TWAU-Net 论文版
输入：doc/论文_基于U-Net与三角窗注意力机制的水葫芦语义分割.md
输出：doc/论文_水利水电技术投稿稿_TWAU-Net.docx

版式沿用 SegFormer 投稿稿的约定：
  - 标题小一黑体 / 作者4号楷体 / 单位5号仿宋
  - 摘要、关键词 5号黑体标签 + 5号楷体内容（结构化【目的】【方法】【结果】【结论】）
  - 中图分类号 + 文献标志码；英文标题/作者/单位/结构化 Abstract/Keywords
  - 引言编号 0，正文自 1 起；一级4号宋体 / 二级5号黑体 / 正文5号宋体单倍行距
  - 图下方中英题名；三线表（表题小五黑体 / 内容六号）；表注六号
  - 公式居中 + 右对齐编号（LaTeX 子集 → Word 文本 + 上下标）
  - 参考文献 6 号（中文者保持中文，不臆造官方英译）
"""
import re, os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = r"D:\chengs\9.project\shuihulu"
MD = os.path.join(ROOT, "twa_unet", "doc", "论文_基于U-Net与三角窗注意力机制的水葫芦语义分割.md")
FIGDIR = os.path.join(ROOT, "twa_unet", "doc", "figs")
OUT = os.path.join(ROOT, "twa_unet", "doc", "论文_水利水电技术投稿稿_TWAU-Net.docx")

# ---------------- 期刊版式常量 ----------------
CN_TITLE = "基于 U-Net 与三角窗注意力机制的水葫芦语义分割方法"
EN_TITLE = ("Semantic Segmentation of Water Hyacinth Based on "
            "U-Net with Triangular Window Attention")
CN_AUTHORS = "程帅1，×××2，×××3，×××4（请补全作者）"
CN_UNIT = "（1. ××××大学 ××××学院，××× 省 ××××；2. ××××单位；3. ……）"
EN_AUTHORS = "CHENG Shuai1, ×××2, ×××3, ×××4 (please complete)"
EN_UNIT = "(1. ×××× University, ××××, China; 2. ……)"

CN_ABS = [
    ("【目的】", "针对无人机航拍河道影像中水葫芦像素级精细识别的需求，提出一种融合三角窗注意力（Triangular Window Attention, TWA）与 U-Net 的语义分割方法 TWAU-Net。"),
    ("【方法】", "与标准卷积核及方形窗口注意力的对称感受野不同，所提 TWA 模块在 M×M 局部窗口内沿主对角线构造下三角或右上三角掩码，使每个位置仅能关注自身左上方或右下方的 token，赋予网络对叶缘、桥墩、船体等方向性纹理的敏感性；在 U-Net 的 4 个跳跃连接处均嵌入 TWA 模块构成 TWAU-Net，并采用类别加权交叉熵与 Dice 损失的加权和缓解类别不均衡。"),
    ("【结果】", "在自建的 2 300 余张太湖流域无人机航拍水葫芦数据集上开展对比、同位等量对照与消融实验，结果表明 TWAU-Net 的水葫芦类 IoU 与整体 mIoU 分别为 0.7807 与 0.8729，较 ResNet34 骨干的 U-Net 提高 18.45% 与 12.42%；在与 TWA 完全相同的 4 个跳跃连接处嵌入同量级的 CBAM 与 Triplet Attention，其水葫芦类 IoU 仍较 TWAU-Net 低 7.60% 与 6.01%；参数量为 27.20 M，单图推理时延为 34.52 ms。"),
    ("【结论】", "消融实验确定下三角方向与窗口尺寸 M=4 为最优组合，三角窗注意力可有效改善对水葫芦叶缘等方向性纹理的像素级分割。"),
]
EN_ABS = [
    ("[Objective]", "To meet the demand of pixel-level fine identification of water hyacinth in UAV images over river channels, a semantic segmentation method named TWAU-Net is proposed by integrating triangular window attention (TWA) into U-Net."),
    ("[Methods]", "Unlike the symmetric receptive fields of standard convolution kernels and square-window attention, the TWA module constructs lower-triangular or upper-right-triangular masks along the main diagonal within an M×M local window, so that each position attends only to tokens in its upper-left or lower-right neighborhood, endowing the network with sensitivity to directional textures such as leaf margins, piers and hulls. TWA modules are embedded at all the four skip connections of U-Net, and a weighted sum of class-weighted cross-entropy and Dice losses is adopted to alleviate class imbalance."),
    ("[Results]", "Comparison, same-position equivalent-control and ablation experiments are conducted on a self-built dataset of more than 2,300 UAV water-hyacinth images collected from rivers and lakes in the Taihu Basin. The results show that TWAU-Net achieves a water-hyacinth IoU of 0.7807 and a mean IoU (mIoU) of 0.8729, 18.45% and 12.42% higher than ResNet34-backed U-Net, respectively; when CBAM and Triplet Attention of similar capacity are embedded at exactly the same four skip connections, their water-hyacinth IoU remains 7.60% and 6.01% lower than that of TWAU-Net, respectively; it has 27.20 M parameters and a single-image inference latency of 34.52 ms."),
    ("[Conclusion]", "Ablation experiments verify that the lower-triangular direction with window size M=4 is the optimal configuration, and the triangular window attention effectively improves the pixel-level segmentation of directional textures such as water hyacinth leaf margins."),
]
CN_KW = "水葫芦；语义分割；U-Net；三角窗注意力；无人机遥感"
EN_KW = "water hyacinth; semantic segmentation; U-Net; triangular window attention; UAV remote sensing"

# 图 / 表 中英题名（键为文件名；编号按正文出现顺序）
FIG_W = {"TWAU-Net_preview.png": 16.0, "TWA_preview.png": 14.0,
         "fig_ablation.png": 13.5, "fig_compare.png": 16.0, "fig_scene_iou.png": 15.0,
         "fig_attention_twa.png": 15.0, "fig_attention_cbam.png": 15.0,
         "fig_attention_triplet.png": 15.0}
FIG_CAP = {
    "TWAU-Net_preview.png": (
        "图1 TWAU-Net 总体结构（编码器为 ResNet34，4 个跳跃连接处各嵌入一个 TWA 模块；TWA 增强后的跳跃特征与解码器上采样特征拼接，最深层特征经上采样汇入 H/16 级拼接；最末两级无拼接，分割头为 1×1 卷积）",
        "Fig.1 Overall architecture of TWAU-Net (the encoder is ResNet34, and a TWA module is embedded at each of the four skip connections; TWA-enhanced skip features are concatenated with the upsampled decoder features, and the deepest features are upsampled and merged into the H/16-level concatenation; the last two stages have no concatenation, and the segmentation head is a 1×1 convolution)"),
    "TWA_preview.png": (
        "图2 三角窗注意力模块（(a) 特征图按 M×M 划分窗口；(b) 窗口内下三角聚合，Query 仅关注自身及左上方 token；(c) 4×4 窗口对应的 16×16 注意力掩码矩阵，深色为 softmax 前置 −∞ 的屏蔽位置；(d) 模块流程：窗口注意力与 FFN 均以双残差（相加后组归一化）形式接入）",
        "Fig.2 Triangular window attention module ((a) feature maps are partitioned into M×M windows; (b) lower-triangular aggregation within a window, where a Query attends only to itself and the tokens above and to its left; (c) the 16×16 attention mask matrix of a 4×4 window, where dark cells denote positions masked by −∞ before softmax; (d) module pipeline: both window attention and FFN are connected in a double-residual form (addition followed by group normalization))"),
    "fig_ablation.png": (
        "图3 消融实验各配置的 mIoU 与水葫芦类 IoU 对比（数据同表 5；横轴第三行为各配置参数量；金框为最终模型 A5：下三角方向、M=4）",
        "Fig.3 mIoU and water-hyacinth IoU of each ablation configuration (data from Table 5; the third row of the x-axis gives the number of parameters of each configuration; the final model A5, i.e. the lower-triangular direction with M=4, is highlighted with a golden frame)"),
    "fig_compare.png": (
        "图4 五种典型场景下各方法的分割结果对比（左起：原图、人工真值、U-Net(ResNet18)、U-Net(ResNet34)、U-Net+CBAM、U-Net+Triplet、DeepLabV3+、TWAU-Net，其中 U-Net+CBAM 与 U-Net+Triplet 为表 2 中 B2、B3 的同位等量 ×4 嵌入基线；水葫芦类以红色高亮，预测图左上角为该样本水葫芦类 IoU，每行 IoU 最高者以金框标出）",
        "Fig.4 Segmentation comparison of different methods under five typical scenes (from left: original image, ground truth, U-Net(ResNet18), U-Net(ResNet34), U-Net+CBAM, U-Net+Triplet, DeepLabV3+, TWAU-Net, where U-Net+CBAM and U-Net+Triplet are the same-position ×4-embedded baselines B2 and B3 in Table 2; the water-hyacinth class is highlighted in red, the upper-left corner of each prediction gives its water-hyacinth IoU, and the best result in each row is framed in gold)"),
    "fig_scene_iou.png": (
        "图5 五种典型场景下各方法水葫芦类 IoU 分组对比（数值与图 4 各面板左上角标注一致；TWAU-Net 在全部 5 个场景均为最优）",
        "Fig.5 Grouped comparison of the water-hyacinth IoU of different methods under the five typical scenes (the values are consistent with those in the upper-left corners of the panels in Fig.4; TWAU-Net is the best in all five scenes)"),
    "fig_attention_twa.png": (
        "图6 TWAU-Net 各尺度注意力热图（自上而下：原图与 s1→s4 四个尺度的注意力；列与图 4 的 5 个验证集场景一一对应；颜色越亮表示注意力越集中）",
        "Fig.6 Multi-scale attention maps of TWAU-Net (from top to bottom: original image and the attention maps at the four scales s1-s4; each column corresponds to one of the five validation scenes in Fig.4; a brighter color indicates more concentrated attention)"),
    "fig_attention_cbam.png": (
        "图7 U-Net+CBAM（×4 同位）各尺度空间门控响应热图（自上而下：原图与 s1→s4 四个尺度的门控响应；列与图 4 的 5 个验证集场景一一对应；颜色越亮表示该位置被门控保留的程度越高）",
        "Fig.7 Multi-scale spatial gating response maps of U-Net+CBAM (×4, same positions) (from top to bottom: original image and the gating responses at the four scales s1-s4; each column corresponds to one of the five validation scenes in Fig.4; a brighter color indicates a higher degree to which the position is retained by the gate)"),
    "fig_attention_triplet.png": (
        "图8 U-Net+Triplet Attn（×4 同位）各尺度空间门控响应热图（行列布局与图 7 一致；颜色越亮表示该位置被门控保留的程度越高）",
        "Fig.8 Multi-scale spatial gating response maps of U-Net+Triplet Attention (×4, same positions) (the row-column layout is the same as Fig.7; a brighter color indicates a higher degree to which the position is retained by the gate)"),
}
# 按正文中出现的先后依次为表 1~5
TAB_CAP = {
    1: ("表1 数据集各类别像素占比统计", "Table 1 Pixel proportion of each class in the dataset"),
    2: ("表2 主结果对比", "Table 2 Main results on the validation set"),
    3: ("表3 代表性方法的逐类 IoU 对比", "Table 3 Per-class IoU comparison of representative methods on the validation set"),
    4: ("表4 TWAU-Net 的行归一化混淆矩阵（%）", "Table 4 Row-normalized confusion matrix of TWAU-Net (%)"),
    5: ("表5 TWAU-Net 消融实验", "Table 5 Ablation experiments of TWAU-Net"),
}

# ---------------- 基础工具 ----------------
def set_run(r, cn="宋体", en="Times New Roman", size=10.5, bold=False, italic=False,
            sub=False, sup=False):
    r.font.name = en
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if sub: r.font.subscript = True
    if sup: r.font.superscript = True
    rPr = r._element.get_or_add_rPr()
    rf = rPr.get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), cn)

def para(doc, align=None, size=10.5, indent_pt=0, space=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(space); pf.space_after = Pt(space)
    if indent_pt: pf.first_line_indent = Pt(indent_pt)
    if align is not None: p.alignment = align
    return p

# ---------------- 数学渲染（LaTeX 子集 → runs） ----------------
SYM = {"alpha":"α","beta":"β","gamma":"γ","delta":"δ","theta":"θ","lambda":"λ",
       "mu":"μ","sigma":"σ","phi":"φ","pi":"π","Omega":"Ω","Sigma":"Σ","Delta":"Δ",
       "cdot":"·","times":"×","div":"÷","sum":"Σ","prod":"∏","in":"∈","infty":"∞",
       "pm":"±","leq":"≤","geq":"≥","ne":"≠","approx":"≈","partial":"∂",
       "rightarrow":"→","to":"→","le":"≤","ge":"≥","ldots":"…"}
BB = {"R":"ℝ","C":"ℂ","N":"ℕ","Z":"ℤ","Q":"ℚ"}
UPRIGHT = {"log","min","max","sin","cos","tan","exp","lim","inf","sup","ConvBlock",
           "TWA","Attn","Up","Concat","CE","Dice"}

def _read_group(tex, i):
    d = 0; j = i
    while j < len(tex):
        if tex[j] == "{": d += 1
        elif tex[j] == "}":
            d -= 1
            if d == 0: return tex[i+1:j], j+1
        j += 1
    return tex[i+1:], len(tex)

def _plain(tex):
    """把不含嵌套上下标的 LaTeX 片段转成普通 Unicode 字符串（用于上/下标）。"""
    out, i, n = [], 0, len(tex)
    while i < n:
        c = tex[i]
        if c == "\\":
            m = re.match(r"\\([a-zA-Z]+)", tex[i:])
            if not m:
                i += 1; continue
            cmd = m.group(1); i += len(m.group(0))
            if cmd in ("mathrm", "text", "mathit", "mathbf"):
                content, j = _read_group(tex, i); i = j
                out.append(content)
            elif cmd == "mathbb":
                content, j = _read_group(tex, i); i = j
                out.append("".join(BB.get(ch, ch) for ch in content))
            elif cmd == "tilde":
                content, j = _read_group(tex, i); i = j
                if content:
                    out.append(content[:-1])
                    out.append(content[-1] + "\u0303")
            elif cmd in ("frac",):
                num, j = _read_group(tex, i)
                den, k = _read_group(tex, j); i = k
                out.append("(" + _plain(num) + ")/(" + _plain(den) + ")")
            elif cmd in SYM:
                out.append(SYM[cmd])
            # 其余命令忽略
            continue
        elif c in "{}":
            i += 1; continue
        elif c == "_" or c == "^":
            i += 1
            if i < n and tex[i] == "{":
                content, j = _read_group(tex, i); i = j
            else:
                content = tex[i]; i += 1
            out.append(_plain(content))
        else:
            out.append(c); i += 1
    return "".join(out)

def render_math(p, tex, size, bold=False, italic=True):
    """把 LaTeX 子集渲染为 runs：普通字符顺次输出，_/^ 内容转为单个上/下标 run。"""
    def add(t, b=bold, it=italic, sub=False, sup=False):
        if t == "":
            return
        r = p.add_run(t)
        set_run(r, cn="Times New Roman", en="Times New Roman", size=size,
                bold=b, italic=it, sub=sub, sup=sup)
    i, n = 0, len(tex)
    acc = ""
    def flush():
        nonlocal acc
        if acc:
            add(acc); acc = ""
    while i < n:
        c = tex[i]
        if c == "\\":
            m = re.match(r"\\([a-zA-Z]+)", tex[i:])
            if not m:
                i += 1; continue
            cmd = m.group(1); i += len(m.group(0))
            if cmd in UPRIGHT:
                flush(); add(cmd, it=False)
            elif cmd in ("mathrm", "text"):
                content, j = _read_group(tex, i); i = j
                flush(); add(content, it=False)
            elif cmd == "mathbf":
                content, j = _read_group(tex, i); i = j
                flush(); add(content, it=False, b=True)
            elif cmd == "mathit":
                content, j = _read_group(tex, i); i = j
                flush(); add(content, it=True)
            elif cmd == "mathbb":
                content, j = _read_group(tex, i); i = j
                flush(); add("".join(BB.get(ch, ch) for ch in content), it=False)
            elif cmd == "tilde":
                content, j = _read_group(tex, i); i = j
                flush()
                if content:
                    add(content[:-1])
                    add(content[-1] + "\u0303")
            elif cmd in ("frac",):
                flush()
                num, j = _read_group(tex, i)
                den, k = _read_group(tex, j); i = k
                add("(" + _plain(num) + ")/(" + _plain(den) + ")", it=True)
            elif cmd in SYM:
                flush(); add(SYM[cmd], it=False)
            else:
                pass  # \left \right \quad \tag \; \, \ { } 等
            continue
        elif c == "_" or c == "^":
            i += 1
            if i < n and tex[i] == "{":
                content, j = _read_group(tex, i); i = j
            else:
                content = tex[i]; i += 1
            flush()
            add(_plain(content), sub=(c == "_"), sup=(c == "^"))
        else:
            acc += c; i += 1
    flush()

# 行内格式：粗体 / 斜体（拉丁学名） / $数学$
INLINE = re.compile(r"(\$[^$]+\$|\*\*[^*]+\*\*|\*[^*]+\*)")

def add_inline(p, text, size=10.5, cn="宋体", en="Times New Roman", bold=False):
    for tok in INLINE.split(text):
        if not tok: continue
        if tok.startswith("$") and tok.endswith("$"):
            render_math(p, tok[1:-1].strip(), size, bold=bold, italic=True)
        elif tok.startswith("**"):
            r = p.add_run(tok[2:-2]); set_run(r, cn=cn, en=en, size=size, bold=True)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = p.add_run(tok[1:-1]); set_run(r, cn=cn, en=en, size=size, italic=True)
        else:
            r = p.add_run(tok); set_run(r, cn=cn, en=en, size=size, bold=bold)

# ---------------- 图 / 表 ----------------
def add_caption(doc, cn_txt, en_txt):
    p1 = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p1.add_run(cn_txt); set_run(r, cn="黑体", en="Times New Roman", size=9)
    p2 = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p2.add_run(en_txt); set_run(r, cn="Times New Roman", en="Times New Roman", size=9)

def add_figure(doc, rel):
    path = os.path.join(ROOT, "twa_unet", "doc", rel.replace("/", os.sep))
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.add_run().add_picture(path, width=Cm(FIG_W.get(os.path.basename(rel), 15.0)))
    cn_cap, en_cap = FIG_CAP[os.path.basename(rel)]
    add_caption(doc, cn_cap, en_cap)

def set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    tb = OxmlElement("w:tcBorders")
    for edge, sz in kw.items():
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
        tb.append(el)
    tcPr.append(tb)

def add_table(doc, rows, tno):
    data = []
    for row in rows:
        cells = [x.strip() for x in row.strip().strip("|").split("|")]
        if cells and re.match(r"^[-:\s]+$", cells[0]):
            continue
        data.append(cells)
    cn_cap, en_cap = TAB_CAP[tno]
    add_caption(doc, cn_cap, en_cap)
    ncol = max(len(r) for r in data)
    tbl = doc.add_table(rows=len(data), cols=ncol)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    for ri, row in enumerate(data):
        for ci in range(ncol):
            cell = tbl.rows[ri].cells[ci]
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            txt = row[ci] if ci < len(row) else ""
            is_hdr = (ri == 0)
            for tok in INLINE.split(txt):
                if not tok: continue
                if tok.startswith("$") and tok.endswith("$"):
                    render_math(cell.paragraphs[0], tok[1:-1].strip(), 7.5,
                                bold=is_hdr, italic=True)
                elif tok.startswith("**"):
                    rr = cell.paragraphs[0].add_run(tok[2:-2])
                    set_run(rr, cn="黑体" if is_hdr else "宋体", en="Times New Roman",
                            size=7.5, bold=True)
                else:
                    rr = cell.paragraphs[0].add_run(tok)
                    set_run(rr, cn="黑体" if is_hdr else "宋体", en="Times New Roman",
                            size=7.5, bold=is_hdr)
    # 三线表
    for c in tbl.rows[0].cells:
        set_cell_border(c, top=12, bottom=4)
    for c in tbl.rows[-1].cells:
        set_cell_border(c, bottom=12)
    return tbl

def add_note(doc, text):
    p = para(doc, indent_pt=15)
    r = p.add_run(text); set_run(r, cn="宋体", en="Times New Roman", size=7.5)

# ---------------- 公式（含 cases 线性化） ----------------
def _expand_cases(tex):
    def rep(m):
        inner = m.group(1)
        rows = [r for r in re.split(r"\\\\", inner)]
        parts = []
        for r in rows:
            r = r.strip()
            if not r: continue
            if "&" in r:
                val, cond = r.split("&", 1)
            else:
                val, cond = r, ""
            val = val.strip(); cond = cond.strip()
            parts.append(val + (("（" + cond + "）") if cond else ""))
        return "\\{ " + "；".join(parts) + " \\}"
    return re.sub(r"\\begin\{cases\}(.*?)\\end\{cases\}", rep, tex, flags=re.S)

def add_equation(doc, latex, tag):
    latex = re.sub(r"\s*\\tag\{\d+\}\s*", "", latex).strip()
    latex = re.sub(r"\\quad|\\qquad", "  ", latex)
    latex = latex.replace("\\,", " ").replace("\\;", " ")
    latex = latex.replace(r"\text{otherwise}", "否则")
    body = _expand_cases(latex)
    lines = [ln for ln in body.splitlines() if ln.strip()]
    for idx, ln in enumerate(lines):
        p = para(doc)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
        render_math(p, ln, 10.5)
        if tag and idx == len(lines) - 1:
            r = p.add_run("\t(%s)" % tag)
            set_run(r, cn="Times New Roman", en="Times New Roman", size=10.5, italic=False)

# ---------------- 章节编号（md 已为 0 引言 / 1-4 正文，原样透传） ----------------
def ren_h1(t):
    return t

def ren_h2(t):
    return t

# ---------------- 主流程 ----------------
def main():
    md = open(MD, encoding="utf-8").read()
    doc = Document()
    for sec in doc.sections:
        sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(2.5)
        sec.left_margin = sec.right_margin = Cm(2.5)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(10.5)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    # ---- 中文标题 / 作者 / 单位 ----
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(CN_TITLE); set_run(r, cn="黑体", size=24, bold=True)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(CN_AUTHORS); set_run(r, cn="楷体", size=14)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(CN_UNIT); set_run(r, cn="仿宋", size=10.5)

    # ---- 中文结构化摘要 / 关键词 ----
    p = para(doc)
    r = p.add_run("摘  要："); set_run(r, cn="黑体", size=10.5, bold=True)
    for tag, seg in CN_ABS:
        r = p.add_run(tag); set_run(r, cn="楷体", size=10.5, bold=True)
        r = p.add_run(seg); set_run(r, cn="楷体", size=10.5)
    p = para(doc)
    r = p.add_run("关键词："); set_run(r, cn="黑体", size=10.5, bold=True)
    r = p.add_run(CN_KW); set_run(r, cn="楷体", size=10.5)
    p = para(doc)
    r = p.add_run("中图分类号：TP391.41　　文献标志码：A"); set_run(r, cn="宋体", size=9)

    # ---- 英文标题 / 作者 / 单位 / Abstract / Keywords ----
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(EN_TITLE); set_run(r, size=14, bold=True)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(EN_AUTHORS); set_run(r, size=10.5)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(EN_UNIT); set_run(r, size=9)
    p = para(doc)
    r = p.add_run("Abstract: "); set_run(r, size=10.5, bold=True)
    for tag, seg in EN_ABS:
        r = p.add_run(tag + " "); set_run(r, size=10.5, bold=True)
        r = p.add_run(seg + " "); set_run(r, size=10.5)
    p = para(doc)
    r = p.add_run("Keywords: "); set_run(r, size=10.5, bold=True)
    r = p.add_run(EN_KW); set_run(r, size=10.5)

    # ---- 基金 / 作者简介 占位 ----
    for note in ["基金项目：××××××××（请填写）",
                 "作者简介：×××（出生年—），男，职称，学位，主要从事×××研究。E-mail：×××",
                 "通信作者：×××（出生年—），男，职称，学位，主要从事×××研究。E-mail：×××"]:
        p = para(doc)
        r = p.add_run(note); set_run(r, cn="宋体", size=9)

    # ---- 正文 ----
    body_src = re.search(r"(## 0 引言\n.*?)\n## 参考文献", md, re.S).group(1)
    body_lines = body_src.split("\n")
    i, n = 0, len(body_lines)
    eq_no = 0
    tno = 0
    while i < n:
        s = body_lines[i].strip()
        if not s or s == "---":
            i += 1; continue
        m1 = re.match(r"^## (.*)$", s)
        if m1:
            p = para(doc)
            r = p.add_run(ren_h1(m1.group(1).strip()))
            set_run(r, cn="宋体", size=14, bold=True)
            p.paragraph_format.space_before = Pt(8)
            i += 1; continue
        m2 = re.match(r"^### (.*)$", s)
        if m2:
            p = para(doc)
            r = p.add_run(ren_h2(m2.group(1).strip()))
            set_run(r, cn="黑体", size=10.5)
            i += 1; continue
        if s == "$$":
            lines = []
            i += 1
            while i < n and body_lines[i].strip() != "$$":
                lines.append(body_lines[i]); i += 1
            i += 1
            eq_no += 1
            add_equation(doc, "\n".join(lines), eq_no)
            continue
        mg = re.match(r"^!\[(.*?)\]\((.*?)\)$", s)
        if mg:
            add_figure(doc, mg.group(2))
            i += 1; continue
        # 跳过独立图/表题（由 add_caption 生成双语题名）
        if re.fullmatch(r"\*\*.+\*\*", s) and (s.startswith("**表") or s.startswith("**图")):
            i += 1; continue
        if s.startswith("|"):
            j = i; rows = []
            while j < n and body_lines[j].strip().startswith("|"):
                rows.append(body_lines[j].strip()); j += 1
            tno += 1
            add_table(doc, rows, tno)
            i = j; continue
        if s.startswith("**注"):
            add_note(doc, s.replace("**", "").strip())
            i += 1; continue
        p = para(doc, indent_pt=21)
        add_inline(p, s, size=10.5, cn="宋体", en="Times New Roman")
        i += 1

    # ---- 参考文献（6 号；中文者保持中文） ----
    ref_src = re.search(r"## 参考文献\n(.*?)(\n---|\Z)", md, re.S).group(1)
    p = para(doc); r = p.add_run("参考文献")
    set_run(r, cn="宋体", size=14, bold=True)
    p.paragraph_format.space_before = Pt(8)
    for line in ref_src.split("\n"):
        line = line.strip()
        if not line: continue
        mm = re.match(r"^\[(\d+)\]", line)
        if not mm: continue
        p = para(doc)
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-18)
        add_inline(p, line, size=7.5, cn="宋体", en="Times New Roman")

    doc.save(OUT)
    print("saved:", OUT)

if __name__ == "__main__":
    main()
